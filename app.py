from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from typing import Dict, List, Tuple

import streamlit as st
from docx import Document
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

APP_NAME = "CareerMatch AI"

# Curated starter taxonomy. Expand this list as you evaluate the app against real job descriptions.
SKILL_TAXONOMY: Dict[str, List[str]] = {
    "Programming Languages": ["python", "java", "javascript", "typescript", "c++", "c#", "sql", "r", "go", "php"],
    "AI / Machine Learning": ["machine learning", "deep learning", "natural language processing", "nlp", "computer vision", "pytorch", "tensorflow", "scikit-learn", "sklearn", "hugging face", "transformers", "generative ai", "llm", "large language models", "rag", "langchain", "embeddings", "prompt engineering", "model evaluation", "feature engineering"],
    "Data": ["pandas", "numpy", "data analysis", "data visualization", "matplotlib", "power bi", "tableau", "excel", "statistics", "etl", "data cleaning"],
    "Web / Backend": ["html", "css", "react", "node.js", "express", "fastapi", "flask", "django", "rest api", "restful api", "api development", "streamlit"],
    "Databases": ["postgresql", "mysql", "mongodb", "sqlite", "redis", "database design", "database management"],
    "Cloud / DevOps": ["aws", "azure", "google cloud", "docker", "kubernetes", "ci/cd", "github actions", "linux", "git", "github", "deployment"],
    "Software Engineering": ["object-oriented programming", "oop", "data structures", "algorithms", "unit testing", "integration testing", "debugging", "agile", "system design", "software development"],
}

SKILL_ALIASES = {
    "sklearn": "scikit-learn",
    "nlp": "natural language processing",
    "llm": "large language models",
    "oop": "object-oriented programming",
    "restful api": "rest api",
}


def normalize_text(text: str) -> str:
    text = text.lower().replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_text(uploaded_file) -> str:
    """Extract text from supported PDF, DOCX, or TXT uploads."""
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(raw))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(BytesIO(raw))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs).strip()
    if name.endswith(".txt"):
        return raw.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")


def find_skills(text: str) -> Dict[str, List[str]]:
    normalized = normalize_text(text)
    found: Dict[str, List[str]] = {}
    for category, skills in SKILL_TAXONOMY.items():
        matches = []
        for skill in skills:
            # Boundaries avoid false positives such as matching 'r' inside unrelated words.
            pattern = rf"(?<![a-z0-9+#.]){re.escape(skill.lower())}(?![a-z0-9+#.])"
            if re.search(pattern, normalized):
                canonical = SKILL_ALIASES.get(skill.lower(), skill)
                if canonical not in matches:
                    matches.append(canonical)
        if matches:
            found[category] = sorted(matches)
    return found


def flatten_skills(grouped: Dict[str, List[str]]) -> List[str]:
    return sorted({skill for skills in grouped.values() for skill in skills})


def compare_resume_to_job(resume_text: str, job_text: str) -> dict:
    resume_skills_grouped = find_skills(resume_text)
    job_skills_grouped = find_skills(job_text)
    resume_skills = set(flatten_skills(resume_skills_grouped))
    job_skills = set(flatten_skills(job_skills_grouped))

    matched = sorted(resume_skills & job_skills)
    missing = sorted(job_skills - resume_skills)
    extra = sorted(resume_skills - job_skills)
    skill_coverage = (len(matched) / len(job_skills) * 100) if job_skills else 0.0

    try:
        matrix = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), max_features=8000).fit_transform([resume_text, job_text])
        text_similarity = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0] * 100)
    except ValueError:
        text_similarity = 0.0

    # A transparent heuristic, not a real ATS score or a hiring probability.
    overall = round((0.7 * skill_coverage) + (0.3 * text_similarity))
    return {
        "resume_skills_grouped": resume_skills_grouped,
        "job_skills_grouped": job_skills_grouped,
        "matched_skills": matched,
        "missing_skills": missing,
        "additional_skills": extra,
        "skill_coverage": round(skill_coverage),
        "text_similarity": round(text_similarity),
        "match_score": overall,
        "job_skill_count": len(job_skills),
    }


def build_recommendations(result: dict, resume_text: str) -> List[str]:
    recommendations: List[str] = []
    missing = result["missing_skills"]
    if missing:
        recommendations.append(
            "Review these job-description skills that were not detected in your resume: "
            + ", ".join(missing[:12])
            + ("." if len(missing) <= 12 else ", and others.")
            + " Add them only if you genuinely have the skill, and support them with evidence."
        )
    else:
        recommendations.append("The skill detector found all listed job-description skills in your resume. Verify that each is supported by clear evidence.")

    text = normalize_text(resume_text)
    if not any(term in text for term in ["project", "projects"]):
        recommendations.append("Add a Projects section with 2–3 relevant projects, your contribution, the tools used, and measurable outcomes where available.")
    if not any(term in text for term in ["experience", "internship", "employment"]):
        recommendations.append("If applicable, add internship, freelance, volunteering, or practical experience. If you are a student, strong academic projects are useful evidence too.")
    if not any(term in text for term in ["github", "portfolio", "linkedin"]):
        recommendations.append("Consider adding relevant portfolio, GitHub, or LinkedIn links so a reviewer can verify your work.")
    if len(resume_text.split()) < 150:
        recommendations.append("The extracted resume text is quite short. Check whether the file is complete or whether it is a scanned PDF that needs OCR.")
    recommendations.append("Tailor the summary and project bullets to the role. Prefer evidence and outcomes over keyword stuffing, and never add skills you do not have.")
    return recommendations


def format_report(result: dict, recommendations: List[str], job_title: str) -> str:
    lines = [
        f"{APP_NAME} — Resume Match Report",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"Target role: {job_title.strip() or 'Not specified'}",
        "",
        f"Estimated match score: {result['match_score']} / 100",
        f"Detected skill coverage: {result['skill_coverage']}%",
        f"Text similarity: {result['text_similarity']}%",
        "",
        "MATCHED SKILLS",
        ", ".join(result["matched_skills"]) or "None detected",
        "",
        "SKILLS NOT DETECTED IN RESUME",
        ", ".join(result["missing_skills"]) or "None detected",
        "",
        "OTHER DETECTED RESUME SKILLS",
        ", ".join(result["additional_skills"]) or "None detected",
        "",
        "RECOMMENDATIONS",
    ]
    lines.extend(f"- {item}" for item in recommendations)
    lines += ["", "Important: This is an experimental keyword and text-similarity estimate, not an actual ATS score, hiring decision, or guarantee of job suitability."]
    return "\n".join(lines)


st.set_page_config(page_title=APP_NAME, page_icon="🎯", layout="wide")
st.markdown(
    """
    <style>
      .block-container {max-width: 1180px; padding-top: 2rem; padding-bottom: 3rem;}
      .hero {padding: 1.5rem 1.7rem; border: 1px solid rgba(128,128,128,.25); border-radius: 18px; margin-bottom: 1.25rem;}
      .muted {opacity: .75; font-size: .95rem;}
      div[data-testid="stMetric"] {border: 1px solid rgba(128,128,128,.22); padding: 1rem; border-radius: 14px;}
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    f"""<div class="hero"><h1>🎯 {APP_NAME}</h1>
    <p>Understand how your resume aligns with a job description, identify skill gaps, and get practical suggestions.</p>
    <p class="muted">Private-by-default demo: analysis runs in this app session. Files are not intentionally saved by this code.</p></div>""",
    unsafe_allow_html=True,
)

with st.sidebar:
    st.header("About this tool")
    st.write("CareerMatch AI compares extracted resume text with a target job description.")
    st.markdown("**What it checks**")
    st.markdown("- Detected skills\n- Skill coverage\n- TF-IDF text similarity\n- Resume improvement suggestions")
    st.info("This is a portfolio MVP. Its skill list is curated and incomplete; results can miss synonyms, context, and skills expressed differently.")
    st.caption("Tip: use a text-based PDF or DOCX. Scanned PDFs may require OCR before upload.")

left, right = st.columns([0.9, 1.1], gap="large")
with left:
    st.subheader("1. Upload your resume")
    resume_file = st.file_uploader("Resume file", type=["pdf", "docx", "txt"], help="Supported: PDF, DOCX, TXT")
    st.subheader("2. Target role")
    job_title = st.text_input("Job title", placeholder="e.g. Junior AI Engineer")
    st.caption("Avoid uploading documents containing sensitive information you do not want processed in this session.")

with right:
    st.subheader("3. Paste the job description")
    job_description = st.text_area(
        "Job description",
        height=270,
        placeholder="Paste the responsibilities, required skills, qualifications, and experience from the job posting...",
        label_visibility="collapsed",
    )
    analyze = st.button("Analyze resume ↗", type="primary", use_container_width=True)

if analyze:
    if not resume_file:
        st.error("Upload a resume file first.")
        st.stop()
    if len(job_description.strip()) < 40:
        st.error("Paste a fuller job description (at least 40 characters) for a more useful comparison.")
        st.stop()
    try:
        resume_text = extract_text(resume_file)
    except Exception as exc:
        st.error(f"Could not read the uploaded file: {exc}")
        st.stop()
    if len(resume_text.strip()) < 30:
        st.error("Very little text could be extracted. If this is a scanned PDF, run OCR first or upload a DOCX/TXT version.")
        st.stop()

    result = compare_resume_to_job(resume_text, job_description)
    recommendations = build_recommendations(result, resume_text)
    st.session_state["analysis"] = {
        "result": result,
        "recommendations": recommendations,
        "job_title": job_title,
        "resume_text": resume_text,
    }

if "analysis" in st.session_state:
    analysis = st.session_state["analysis"]
    result = analysis["result"]
    recommendations = analysis["recommendations"]
    st.divider()
    st.subheader("Your match overview")
    st.caption("These are transparent heuristic indicators, not a real ATS score or a prediction of whether you will get hired.")
    m1, m2, m3 = st.columns(3)
    m1.metric("Estimated match", f"{result['match_score']} / 100")
    m2.metric("Detected skill coverage", f"{result['skill_coverage']}%")
    m3.metric("Text similarity", f"{result['text_similarity']}%")
    st.progress(min(result["match_score"], 100) / 100)

    tab1, tab2, tab3 = st.tabs(["Skill comparison", "Recommendations", "Detected skills"])
    with tab1:
        a, b = st.columns(2)
        with a:
            st.markdown("#### ✅ Skills found in both")
            if result["matched_skills"]:
                for skill in result["matched_skills"]:
                    st.markdown(f"- {skill}")
            else:
                st.write("No matching skills detected from the current taxonomy.")
        with b:
            st.markdown("#### 🔎 Skills not detected in resume")
            if result["missing_skills"]:
                for skill in result["missing_skills"]:
                    st.markdown(f"- {skill}")
            else:
                st.write("No job-description skills were missing from the detected list.")
        if not result["job_skill_count"]:
            st.warning("No skills from the built-in taxonomy were detected in the job description. The taxonomy may need to be expanded for this role.")
    with tab2:
        for idx, recommendation in enumerate(recommendations, start=1):
            st.markdown(f"**{idx}.** {recommendation}")
    with tab3:
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("#### Resume skills")
            if result["resume_skills_grouped"]:
                for category, skills in result["resume_skills_grouped"].items():
                    st.markdown(f"**{category}**: {', '.join(skills)}")
            else:
                st.write("No skills detected. Check the extracted text and the built-in taxonomy.")
        with c2:
            st.markdown("#### Job-description skills")
            if result["job_skills_grouped"]:
                for category, skills in result["job_skills_grouped"].items():
                    st.markdown(f"**{category}**: {', '.join(skills)}")
            else:
                st.write("No skills detected.")

    report = format_report(result, recommendations, analysis["job_title"])
    st.download_button("Download report (.txt)", data=report, file_name="careermatch_report.txt", mime="text/plain")
    with st.expander("Preview extracted resume text"):
        st.text(analysis["resume_text"][:12000])

st.divider()
st.caption("CareerMatch AI is an educational portfolio project. It does not make hiring decisions, verify qualifications, or guarantee that a resume will pass an employer's screening process.")
